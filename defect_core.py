"""Core logic for the defect coding app: history storage, reading uploads,
auto-coding with Claude, the monthly summary, and Excel exports.

Nothing in this file depends on Streamlit, so it can be reused from a script.
"""
from __future__ import annotations

import io
import json
import math
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

DATA_DIR = Path(os.environ.get("DEFECT_DATA_DIR", Path(__file__).parent / "data"))
HISTORY_FILE = DATA_DIR / "history.json"
SETTINGS_FILE = DATA_DIR / "settings.json"

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
DATA_HEADERS = ["Part Number", "Date", "Program", "Defect Description", "defect code"]

RED = "F4B6B6"
YELLOW = "FFE699"
HEADER_FILL = "D9E1F2"


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "").strip().lower())


def month_label(key: str) -> str:
    year, month = key.split("-")
    return f"{MONTHS[int(month) - 1]} {year}"


def _atomic_write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=1, ensure_ascii=False)
    os.replace(tmp, path)


# --------------------------------------------------------------------------
# History and settings
# --------------------------------------------------------------------------
def load_history() -> dict:
    """Returns {"YYYY-MM": {"rows": [...], "saved_by": str, "saved_at": str, "source": str}}."""
    if not HISTORY_FILE.exists():
        return {}
    return json.loads(HISTORY_FILE.read_text(encoding="utf-8")).get("months", {})


def save_history(months: dict) -> None:
    _atomic_write_json(HISTORY_FILE, {"months": dict(sorted(months.items()))})


def save_month(key: str, rows: list[dict], saved_by: str, source: str) -> None:
    months = load_history()  # re-read so another user's save isn't lost
    months[key] = {"rows": rows, "saved_by": saved_by or "Unknown",
                   "saved_at": datetime.now().strftime("%Y-%m-%d %H:%M"), "source": source}
    save_history(months)


def delete_month(key: str) -> None:
    months = load_history()
    months.pop(key, None)
    save_history(months)


def load_settings() -> dict:
    if not SETTINGS_FILE.exists():
        return {"codes": [], "rules": []}
    data = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
    return {"codes": list(data.get("codes", [])), "rules": list(data.get("rules", []))}


def save_settings(codes: list[str], rules: list[str]) -> None:
    codes = sorted({c.strip() for c in codes if c.strip()}, key=str.lower)
    _atomic_write_json(SETTINGS_FILE, {"codes": codes, "rules": [r.strip() for r in rules if r.strip()]})


def history_rows(months: dict) -> list[dict]:
    return [r for key in sorted(months) for r in months[key]["rows"]]


def all_codes(months: dict, settings: dict) -> list[str]:
    codes = set(settings["codes"])
    codes.update(r["code"] for r in history_rows(months) if r.get("code"))
    return sorted(codes, key=str.lower)


# --------------------------------------------------------------------------
# Reading an uploaded month
# --------------------------------------------------------------------------
def _fmt_date(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    return str(value).strip()


def _cell(value) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    return str(value).strip()


def read_upload(file_bytes: bytes, filename: str) -> dict[str, list[dict]]:
    """Reads every worksheet. Columns A-D are part, date, program, description.
    If the file was already coded, columns E-G (code, confidence, notes) are kept too."""
    sheets: dict[str, list[dict]] = {}
    if filename.lower().endswith(".csv"):
        frames = {"Sheet1": pd.read_csv(io.BytesIO(file_bytes), header=None, dtype=object)}
        raw = {name: df.values.tolist() for name, df in frames.items()}
    else:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        raw = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    for name, table in raw.items():
        rows = []
        for i, r in enumerate(table):
            r = list(r) + [None] * (8 - len(r))
            desc = _cell(r[3])
            if not desc:
                continue
            if i == 0 and ("desc" in desc.lower() or "part" in _cell(r[0]).lower()):
                continue  # header row
            rows.append({"pn": _cell(r[0]), "date": _fmt_date(r[1]), "prog": _cell(r[2]), "desc": desc,
                         "file_code": _cell(r[4]), "file_conf": _cell(r[5]), "file_note": _cell(r[6]),
                         "file_decision": _cell(r[7])})
        sheets[name] = rows
    return sheets


def guess_month(rows: list[dict]) -> str:
    counts: dict[str, int] = {}
    for r in rows:
        m = re.match(r"^(\d{4})-(\d{2})", r["date"])
        if m:
            counts[m.group(0)] = counts.get(m.group(0), 0) + 1
    if counts:
        return max(counts, key=counts.get)
    return date.today().strftime("%Y-%m")


# --------------------------------------------------------------------------
# The working draft for one month
# --------------------------------------------------------------------------
# Each draft row: pn, date, prog, desc, code, conf ("High"|"Medium"|"None"|""),
# note, decision ("Auto"|"Accepted"|"Excluded"|"Needs review"|"Changed"|"")

def new_draft_rows(rows: list[dict], codes: list[str]) -> list[dict]:
    """Creates draft rows. Codes already in the file (columns E-G) are loaded."""
    conf_map = {"high": "High", "accepted": "High", "medium": "Medium", "low": "Medium", "none": "None"}
    out = []
    for r in rows:
        row = {"pn": r["pn"], "date": r["date"], "prog": r["prog"], "desc": r["desc"],
               "code": "", "conf": "", "note": "", "decision": ""}
        file_code, file_conf = r.get("file_code", ""), r.get("file_conf", "")
        if file_code or file_conf:
            match = next((c for c in codes if norm(c) == norm(file_code)), "")
            conf = conf_map.get(norm(file_conf), "High" if match else "None")
            note = r.get("file_note", "")
            if file_code and not match and norm(file_code) != "no match":
                note = f'File had "{file_code}", which isn\'t in the code list. {note}'.strip()
            if not match:
                conf = "None"
            decision = "Auto" if conf == "High" else "Needs review"
            saved = r.get("file_decision", "")
            if saved in ("Accepted", "Changed", "Excluded") and (match or saved == "Excluded"):
                decision = saved  # keeps review decisions from a re-uploaded working file
            row.update(code=match, conf=conf, note=note, decision=decision)
        out.append(row)
    return out


def code_exact_matches(draft: list[dict], months: dict) -> int:
    """Codes rows whose description matches a past one word for word. Returns how many."""
    exact = {norm(r["desc"]): r["code"] for r in history_rows(months)}
    n = 0
    for row in draft:
        if row["conf"]:
            continue
        hit = exact.get(norm(row["desc"]))
        if hit:
            row.update(code=hit, conf="High", note="Same description as a past month", decision="Auto")
            n += 1
    return n


def is_flagged(row: dict) -> bool:
    return row["conf"] in ("Medium", "None")


def save_problem(draft: list[dict]) -> str | None:
    """Returns why the month can't be saved yet, or None if it can."""
    if not draft:
        return "Upload a month first."
    uncoded = sum(1 for r in draft if not r["conf"])
    if uncoded:
        return f"{uncoded} rows aren't coded yet."
    pending = sum(1 for r in draft if r["decision"] == "Needs review")
    if pending:
        return f"{pending} flagged rows still need a decision (Accepted, Changed, or Excluded)."
    if any(r["decision"] != "Excluded" and not r["code"] for r in draft):
        return "Every row you keep needs a defect code."
    if all(r["decision"] == "Excluded" for r in draft):
        return "Every row is excluded."
    return None


def rows_to_save(draft: list[dict]) -> list[dict]:
    return [{"pn": r["pn"], "date": r["date"], "prog": r["prog"], "desc": r["desc"], "code": r["code"]}
            for r in draft if r["decision"] != "Excluded"]


# --------------------------------------------------------------------------
# Auto-coding with Claude
# --------------------------------------------------------------------------
STOPWORDS = set("s/b sb should actual the and are was has have that with from not but its it's there their "
                "this into onto for any all one both out off been being then than they them get got just "
                "way more like also only".split())


def _tokens(text: str) -> set[str]:
    words = re.sub(r"[^a-z0-9+\-/ ]", " ", norm(text)).split()
    return {w for w in words if len(w) > 2 and w not in STOPWORDS}


def build_prompt(batch: list[tuple[int, dict]], months: dict, settings: dict, codes: list[str]) -> str:
    """Builds one prompt for a batch of rows, using the most similar past examples."""
    unique = {}
    for r in history_rows(months):
        unique.setdefault(norm(r["desc"]), r)
    hist = [(r, _tokens(r["desc"])) for r in unique.values()]

    picked: dict[str, dict] = {}
    for _, row in batch:
        t = _tokens(row["desc"])
        scored = sorted(((len(t & ht) / math.sqrt(len(ht) + 1), r) for r, ht in hist if t & ht),
                        key=lambda x: -x[0])[:6]
        for _, r in scored:
            picked[norm(r["desc"])] = r
    per_code: dict[str, list[dict]] = {}
    for r, _ in hist:
        per_code.setdefault(r["code"], []).append(r)
    for c in codes:  # at least two examples of every code, so rare codes are visible
        for r in per_code.get(c, [])[:2]:
            picked[norm(r["desc"])] = r

    examples = "\n".join(f'{re.sub(r"\s+", " ", r["desc"])[:240]}  =>  {r["code"]}'
                         for r in list(picked.values())[:400])
    rules = "\n".join(f"{i}. {r}" for i, r in enumerate(settings["rules"], 1)) or "(none)"
    items = "\n".join(f'{i}: {re.sub(r"\s+", " ", row["desc"])[:600]}' for i, row in batch)
    code_list = "\n".join(f"- {c}" for c in codes)

    return f"""You are helping a manufacturing quality team assign a defect code to each defect description. Descriptions use the pattern "s/b: <what it should be> Actual: <what was found>".

ALLOWED DEFECT CODES (use one exactly as written):
{code_list}

CODING RULES (these override the examples):
{rules}

EXAMPLES FROM PAST MONTHS (description => code the team assigned):
{examples}

NEW DEFECTS TO CODE:
{items}

For each new defect choose the single best allowed code. Follow how the team coded similar examples.
confidence:
- "High": clearly matches one code and fits the examples and rules.
- "Medium": a code fits but the description is ambiguous, could fit two codes, or is a pattern not seen before. The team will review it.
- "None": no allowed code fits, it is too vague to classify, or it is not a real defect description. Set code to null.
note: empty for High. For Medium or None give a short reason (under 25 words); for None name the nearest code if one is close.

Reply with only a JSON array, one object per new defect, in the same order:
[{{"i":0,"code":"Torque Defect","confidence":"High","note":""}}]"""


def parse_json_array(text: str) -> list:
    text = text.strip()
    fence = re.search(r"```(?:json)?\s*(.*?)```", text, re.S)
    if fence:
        text = fence.group(1).strip()
    start, end = text.find("["), text.rfind("]")
    if start == -1 or end == -1:
        raise ValueError("No JSON array in Claude's reply")
    data = json.loads(text[start:end + 1])
    if not isinstance(data, list):
        raise ValueError("Claude's reply wasn't a list")
    return data


def apply_suggestions(batch: list[tuple[int, dict]], suggestions: list, codes: list[str]) -> None:
    by_index = {}
    for s in suggestions:
        if isinstance(s, dict) and "i" in s:
            try:
                by_index[int(s["i"])] = s
            except (TypeError, ValueError):
                pass
    for i, row in batch:
        s = by_index.get(i)
        if not s:
            continue  # left uncoded; the next run picks it up
        code = s.get("code") if isinstance(s.get("code"), str) else ""
        conf = s.get("confidence") if s.get("confidence") in ("High", "Medium", "None") else "Medium"
        note = s.get("note") if isinstance(s.get("note"), str) else ""
        if code:
            match = next((c for c in codes if norm(c) == norm(code)), "")
            if not match:
                note = f'Suggested "{code}", which isn\'t in the code list. {note}'.strip()
                conf = "None"
            code = match
        if not code:
            conf = "None"
        row.update(code=code, conf=conf, note=note.strip(),
                   decision="Auto" if conf == "High" else "Needs review")


def auto_code(draft: list[dict], months: dict, settings: dict, client, model: str,
              batch_size: int = 20, progress=None) -> None:
    """Codes every uncoded row in place. `client` is an anthropic.Anthropic().
    `progress(done, total)` is called after each batch."""
    codes = all_codes(months, settings)
    todo = [(i, r) for i, r in enumerate(draft) if not r["conf"]]
    batches = [todo[k:k + batch_size] for k in range(0, len(todo), batch_size)]
    for n, batch in enumerate(batches, 1):
        prompt = build_prompt(batch, months, settings, codes)
        resp = client.messages.create(model=model, max_tokens=4000,
                                      messages=[{"role": "user", "content": prompt}])
        text = "".join(block.text for block in resp.content if getattr(block, "type", "") == "text")
        apply_suggestions(batch, parse_json_array(text), codes)
        if progress:
            progress(n, len(batches))


# --------------------------------------------------------------------------
# Monthly summary
# --------------------------------------------------------------------------
def column_names(months: dict) -> dict[str, str]:
    """Month column headers: "January" when all months share a year, else "January 2025"."""
    years = {k[:4] for k in months}
    return {k: (MONTHS[int(k[5:7]) - 1] if len(years) == 1 else month_label(k)) for k in sorted(months)}


def summary_table(months: dict, settings: dict) -> pd.DataFrame:
    codes = all_codes(months, settings)
    names = column_names(months)
    data = {names[k]: [sum(1 for r in months[k]["rows"] if r["code"] == c) for c in codes]
            for k in sorted(months)}
    df = pd.DataFrame(data, index=pd.Index(codes, name="defect code"))
    df["Total"] = df.sum(axis=1)
    return df


# --------------------------------------------------------------------------
# Excel exports
# --------------------------------------------------------------------------
_ARIAL = Font(name="Arial", size=10)
_ARIAL_BOLD = Font(name="Arial", size=10, bold=True)
_THIN = Side(style="thin", color="BFBFBF")


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.font = _ARIAL_BOLD
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.border = Border(bottom=_THIN)
    ws.freeze_panes = "A2"


def _data_sheet(wb: Workbook, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(DATA_HEADERS)
    for r in rows:
        ws.append([r["pn"], r["date"], r["prog"], r["desc"], r["code"]])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = _ARIAL
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 4)
    for col, width in zip("ABCDE", (14, 12, 10, 90, 44)):
        ws.column_dimensions[col].width = width
    _style_header(ws, 5)


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def coded_month_xlsx(draft: list[dict], key: str) -> bytes:
    """The month's working file: code, confidence, notes, decision, red/yellow highlights."""
    wb = Workbook()
    ws = wb.active
    ws.title = MONTHS[int(key[5:7]) - 1]
    ws.append(DATA_HEADERS + ["Confidence", "Notes", "Decision"])
    for r in draft:
        code = (r["code"] or "NO MATCH") if r["conf"] else ""
        ws.append([r["pn"], r["date"], r["prog"], r["desc"], code, r["conf"], r["note"], r["decision"]])
    for i, r in enumerate(draft, start=2):
        fill = RED if r["conf"] == "None" else YELLOW if r["conf"] == "Medium" else None
        for c in range(1, 9):
            cell = ws.cell(i, c)
            cell.font = _ARIAL
            cell.alignment = Alignment(vertical="top", wrap_text=c in (4, 7))
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
    for col, width in zip("ABCDEFGH", (14, 12, 10, 80, 44, 12, 60, 14)):
        ws.column_dimensions[col].width = width
    _style_header(ws, 8)
    return _to_bytes(wb)


def summary_xlsx(months: dict, settings: dict) -> bytes:
    """Summary tab (codes in column A, one column per month, COUNTIF formulas)
    plus one data tab per month, the same layout as Defect_Code_Monthly_Summary.xlsx."""
    codes = all_codes(months, settings)
    names = column_names(months)
    keys = sorted(months)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["defect code"] + [names[k] for k in keys])
    for i, code in enumerate(codes, start=2):
        ws.cell(i, 1, code).font = _ARIAL
        for j, k in enumerate(keys, start=2):
            sheet = f"{names[k]} Data"[:31]
            last = len(months[k]["rows"]) + 1
            cell = ws.cell(i, j, f"=COUNTIF('{sheet}'!$E$2:$E${last},$A{i})")
            cell.font = _ARIAL
            cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=1, max_row=len(codes) + 1):
        for cell in row:
            cell.border = Border(bottom=_THIN)
    ws.column_dimensions["A"].width = 46
    for j in range(2, len(keys) + 2):
        ws.cell(1, j).alignment = Alignment(horizontal="center")
        ws.column_dimensions[ws.cell(1, j).column_letter].width = 11
    _style_header(ws, len(keys) + 1)
    ws.freeze_panes = "B2"
    for k in keys:
        _data_sheet(wb, f"{names[k]} Data", months[k]["rows"])
    return _to_bytes(wb)


def history_xlsx(months: dict) -> bytes:
    """Full history, newest month first, one tab per month."""
    names = column_names(months)
    wb = Workbook()
    wb.remove(wb.active)
    for k in sorted(months, reverse=True):
        _data_sheet(wb, names[k], months[k]["rows"])
    return _to_bytes(wb)


def import_history_xlsx(file_bytes: bytes, year: int) -> dict[str, list[dict]]:
    """Reads a history workbook (one tab per month, tab named after the month) into
    {"YYYY-MM": rows}. Used to load an existing workbook into the app."""
    out = {}
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    for ws in wb.worksheets:
        name = ws.title.strip().lower()
        idx = next((i for i, m in enumerate(MONTHS) if name.startswith(m.lower()[:3])), None)
        if idx is None:
            continue
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            r = list(r) + [None] * (5 - len(r))
            if i == 0 or not _cell(r[3]) or not _cell(r[4]):
                continue
            rows.append({"pn": _cell(r[0]), "date": _fmt_date(r[1]), "prog": _cell(r[2]),
                         "desc": _cell(r[3]), "code": _cell(r[4])})
        if rows:
            out[f"{year}-{idx + 1:02d}"] = rows
    return out
